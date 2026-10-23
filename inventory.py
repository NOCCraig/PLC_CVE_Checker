"""Load the PLC inventory from an Excel (.xlsx/.xlsm) or CSV file."""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from .models import PLC

# Header names we recognise (compared after lower-casing and dropping
# anything that isn't a letter or digit).
HEADER_ALIASES: dict[str, list[str]] = {
    "vendor": ["vendor", "manufacturer", "make", "brand", "mfr", "mfg", "oem"],
    "model": [
        "model", "modelnumber", "modelno", "partnumber", "partno", "part",
        "catalognumber", "catalogno", "catalog", "catno", "plcmodel", "plc",
        "product", "productname", "type", "family",
    ],
    "location": ["location", "site", "plant", "area", "facility", "building", "line"],
    "owner": ["owner", "responsible", "contact", "contactname", "department", "dept", "team"],
    # Optional: when present, each PLC's firmware is checked against affected versions.
    "firmware": ["firmware", "firmwareversion", "firmwarerev", "fw", "fwversion", "fwrev",
                 "firmwarerevision", "version", "revision"],
}

REQUIRED = ("vendor", "model")


class InventoryError(Exception):
    pass


@dataclass
class Inventory:
    plcs: list[PLC]
    warnings: list[str] = field(default_factory=list)
    columns: dict[str, str] = field(default_factory=dict)  # field -> header text used


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _read_rows(path: Path, sheet: str | None) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet:
                if sheet not in wb.sheetnames:
                    raise InventoryError(
                        f"Sheet '{sheet}' not found. Sheets in file: {', '.join(wb.sheetnames)}"
                    )
                ws = wb[sheet]
            else:
                ws = wb.worksheets[0]
            return [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    if suffix in (".csv", ".txt"):
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [[_cell(c) for c in row] for row in csv.reader(f)]
    raise InventoryError(f"Unsupported inventory file type '{suffix}'. Use .xlsx, .xlsm or .csv.")


def _map_columns(header: list[str], overrides: dict[str, str]) -> dict[str, int]:
    """Return field -> column index for one header row."""
    normed = [_norm(h) for h in header]
    mapping: dict[str, int] = {}
    for fld in HEADER_ALIASES:
        wanted = overrides.get(fld)
        if wanted:
            if _norm(wanted) in normed:
                mapping[fld] = normed.index(_norm(wanted))
            continue
        for alias in HEADER_ALIASES[fld]:
            if alias in normed:
                mapping[fld] = normed.index(alias)
                break
    return mapping


def load_inventory(
    path: str | Path,
    sheet: str | None = None,
    column_overrides: dict[str, str] | None = None,
) -> Inventory:
    path = Path(path)
    if not path.exists():
        raise InventoryError(f"Inventory file not found: {path}")
    overrides = {k: v for k, v in (column_overrides or {}).items() if v}
    rows = _read_rows(path, sheet)
    if not rows:
        raise InventoryError("The inventory file is empty.")

    header_idx, mapping = None, {}
    for i, row in enumerate(rows[:15]):
        m = _map_columns(row, overrides)
        if all(r in m for r in REQUIRED):
            header_idx, mapping = i, m
            break
    if header_idx is None:
        first = [h for h in rows[0] if h]
        raise InventoryError(
            "Couldn't find a header row with both a vendor and a model column. "
            f"Columns in the first row: {first}. "
            "Rename them (e.g. 'Vendor', 'Model') or use --vendor-col / --model-col."
        )

    header = rows[header_idx]
    inv = Inventory(plcs=[], columns={f: header[i] for f, i in mapping.items()})

    def get(row: list[str], fld: str) -> str:
        i = mapping.get(fld)
        return row[i] if i is not None and i < len(row) else ""

    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(row):
            continue
        vendor, model = get(row, "vendor"), get(row, "model")
        if not vendor or not model:
            inv.warnings.append(f"Row {offset}: missing vendor or model, skipped.")
            continue
        inv.plcs.append(
            PLC(vendor=vendor, model=model, location=get(row, "location"),
                owner=get(row, "owner"), row=offset, firmware=get(row, "firmware"))
        )
    if not inv.plcs:
        raise InventoryError("No usable PLC rows found below the header.")
    return inv
